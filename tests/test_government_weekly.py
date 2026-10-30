from __future__ import annotations

import csv
import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from src.government_weekly import build_week, previous_week


class GovernmentWeeklyTests(unittest.TestCase):
    def write_day(self, root: Path, day: date, *, notice: bool = False, attachment: bool = False):
        day_root = root / day.isoformat()
        day_root.mkdir()
        folder = Path("政策正文") / "国家能源局—通知" / "储能公告_abc"
        file = Path("附件") / "国家能源局—通知" / "储能公告_abc" / "储能公告_附件1.pdf"
        notices = []
        if notice:
            (day_root / folder).mkdir(parents=True)
            (day_root / folder / "正文.txt").write_text(f"{day} 正文", encoding="utf-8")
            items = []
            if attachment:
                (day_root / file).parent.mkdir(parents=True)
                (day_root / file).write_bytes(f"%PDF-{day}".encode())
                items.append({"status": "downloaded", "file": str(file), "url": "https://www.nea.gov.cn/a.pdf", "bytes": (day_root / file).stat().st_size})
            notices = [{"title": "储能公告", "published_at": day.isoformat(), "url": f"https://www.nea.gov.cn/{day}",
                        "folder": str(folder), "attachment_folder": str(file.parent),
                        "content_status": "full_text", "attachments": items, "errors": []}]
        index = {"target_date": day.isoformat(), "sources": [
            {"source": "国家能源局—通知", "status": "ok", "listing_count": len(notices), "notices": notices, "errors": []},
            {"source": "自然资源部—通知公告", "status": "listing_failed", "listing_count": 0,
             "notices": [], "errors": ["blocked by robots.txt"]},
        ]}
        (day_root / "索引.json").write_text(json.dumps(index, ensure_ascii=False), encoding="utf-8")
        with (day_root / "国内政策清单.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=("渠道", "发布日期", "公告名称", "原文链接", "正文状态", "公告文件夹", "政策分类路径", "政策第5级", "政策分类状态"))
            writer.writeheader()
            for item in notices:
                writer.writerow({"渠道": "国家能源局—通知", "发布日期": day.isoformat(),
                                 "公告名称": item["title"], "原文链接": item["url"],
                                 "正文状态": "full_text", "公告文件夹": str(folder),
                                 "政策分类路径": "政策 > 监管落实型政策 > 正外部性 > 消费端优惠 > 初次购置补贴",
                                 "政策第5级": "初次购置补贴", "政策分类状态": "classified"})

    def test_previous_week_and_monday_validation(self):
        self.assertEqual(previous_week(date(2026, 10, 19)), date(2026, 10, 12))
        self.assertEqual(previous_week(date(2026, 10, 21)), date(2026, 10, 12))
        with tempfile.TemporaryDirectory() as temp, self.assertRaisesRegex(ValueError, "周一"):
            build_week(Path(temp), date(2026, 10, 13))

    def test_combines_seven_days_and_rebuilds_without_duplicate_files(self):
        start = date(2026, 10, 12)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for offset in range(7):
                self.write_day(root, start + timedelta(days=offset), notice=offset in (0, 1), attachment=offset in (0, 1))
            target, info = build_week(root, start)
            self.assertEqual(info["policy_count"], 2)
            self.assertEqual(info["attachment_count"], 2)
            self.assertEqual(info["missing_days"], [])
            self.assertEqual(info["missing_files"], [])
            self.assertIn("2026-W42_2026-10-12_2026-10-18", str(target))
            for offset in (0, 1):
                day = start + timedelta(days=offset)
                attachment = target / "附件" / day.isoformat() / "国家能源局—通知" / "储能公告_abc" / "储能公告_附件1.pdf"
                self.assertEqual(attachment.read_bytes(), f"%PDF-{day}".encode())
                metadata = json.loads((target / "政策正文" / day.isoformat() / "国家能源局—通知" / "储能公告_abc" / "公告.json").read_text())
                self.assertEqual(metadata["attachments"][0]["file"], str(attachment.relative_to(target)))
            book = load_workbook(target / "国内政策周汇总.xlsx")
            self.assertEqual(book.sheetnames, ["政策汇总", "附件明细", "渠道状态"])
            self.assertEqual(book["政策汇总"].max_row, 3)
            self.assertEqual(book["附件明细"].max_row, 3)
            self.assertEqual(book["渠道状态"].max_row, 15)
            self.assertEqual(book["政策汇总"]["G2"].value, 1)
            self.assertEqual(book["政策汇总"]["A2"].value.date(), start)
            self.assertEqual(book["政策汇总"]["J2"].value, "政策 > 监管落实型政策 > 正外部性 > 消费端优惠 > 初次购置补贴")
            self.assertEqual(book["政策汇总"]["O2"].value, "初次购置补贴")
            self.assertTrue(book["附件明细"]["E2"].hyperlink)
            target_again, info_again = build_week(root, start)
            self.assertEqual(target_again, target)
            self.assertEqual(info_again["attachment_count"], 2)
            self.assertEqual(len(list((root / "weekly").iterdir())), 1)

    def test_missing_daily_archive_is_visible_in_workbook(self):
        start = date(2026, 10, 12)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for offset in range(6):
                self.write_day(root, start + timedelta(days=offset))
            target, info = build_week(root, start)
            self.assertEqual(info["missing_days"], ["2026-10-18"])
            book = load_workbook(target / "国内政策周汇总.xlsx")
            values = [cell.value for row in book["渠道状态"].iter_rows() for cell in row]
            self.assertIn("daily_missing", values)
            self.assertTrue((target / "附件" / "说明.txt").is_file())

    def test_missing_downloaded_attachment_is_not_reported_as_saved(self):
        start = date(2026, 10, 12)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for offset in range(7):
                self.write_day(root, start + timedelta(days=offset), notice=offset == 0, attachment=offset == 0)
            attachment = next((root / start.isoformat() / "附件").rglob("*.pdf"))
            attachment.unlink()
            target, info = build_week(root, start)
            self.assertEqual(info["attachment_count"], 0)
            self.assertEqual(len(info["missing_files"]), 1)
            book = load_workbook(target / "国内政策周汇总.xlsx")
            self.assertEqual(book["附件明细"]["H2"].value, "file_missing")
            self.assertEqual(book["政策汇总"]["G2"].value, 0)

    def test_daily_index_and_list_must_have_same_notices(self):
        start = date(2026, 10, 12)
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for offset in range(7):
                self.write_day(root, start + timedelta(days=offset), notice=offset == 0)
            (root / start.isoformat() / "国内政策清单.csv").write_text(
                "渠道,发布日期,公告名称,原文链接,正文状态,公告文件夹\n", encoding="utf-8"
            )
            with self.assertRaisesRegex(ValueError, "条目不一致"):
                build_week(root, start)


if __name__ == "__main__":
    unittest.main()
