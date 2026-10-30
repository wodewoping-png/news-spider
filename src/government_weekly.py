"""Build a self-contained Monday policy archive from the preceding seven daily archives."""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TIMEZONE = ZoneInfo("Asia/Shanghai")
POLICY_CLASSIFICATION_COLUMNS = ("政策分类路径", "政策第1级", "政策第2级", "政策第3级", "政策第4级", "政策第5级", "政策分类状态", "政策分类依据", "政策分类明细")
POLICY_COLUMNS = ("公告日期", "渠道", "公告名称", "原文链接", "正文状态", "正文文件", "附件数量", "附件文件夹", "附件下载") + POLICY_CLASSIFICATION_COLUMNS
ATTACHMENT_COLUMNS = ("公告日期", "渠道", "公告名称", "附件名称", "周归档路径", "原始地址", "文件大小（字节）", "下载状态", "说明")
STATUS_COLUMNS = ("归档日期", "渠道", "抓取状态", "列表条目数", "公告数", "错误说明")


def previous_week(today: date) -> date:
    return today - timedelta(days=today.weekday() + 7)


def checked_file(root: Path, relative: str, prefix: str) -> Path:
    path = Path(relative)
    if path.is_absolute() or ".." in path.parts or not path.parts or path.parts[0] != prefix:
        raise ValueError(f"非法{prefix}路径: {relative}")
    full = root / path
    if not full.is_file() or full.is_symlink() or not full.resolve().is_relative_to(root.resolve()):
        raise FileNotFoundError(f"归档文件不存在: {full}")
    return full


def save_workbook(destination: Path, policies: list[dict], attachments: list[dict], statuses: list[dict]) -> None:
    book = Workbook()
    book.remove(book.active)
    specs = (
        ("政策汇总", POLICY_COLUMNS, policies, {"原文链接", "正文文件"}),
        ("附件明细", ATTACHMENT_COLUMNS, attachments, {"周归档路径", "原始地址"}),
        ("渠道状态", STATUS_COLUMNS, statuses, set()),
    )
    widths = {
        "政策汇总": (15, 31, 65, 64, 16, 72, 12, 72, 90, 70, 15, 26, 24, 25, 28, 20, 60, 90),
        "附件明细": (15, 31, 65, 65, 90, 64, 18, 15, 44),
        "渠道状态": (15, 31, 18, 16, 13, 54),
    }
    for name, columns, records, linked in specs:
        sheet = book.create_sheet(name)
        sheet.append(columns)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for row in records:
            sheet.append([row.get(column, "") for column in columns])
            for col, column in enumerate(columns, 1):
                cell = sheet.cell(sheet.max_row, col)
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="center")
                if column in ("公告日期", "归档日期"):
                    cell.number_format = "yyyy-mm-dd"
                if column in linked and cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = Font(name="Microsoft YaHei", size=10, color="0563C1", underline="single")
        for i, width in enumerate(widths[name], 1):
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
        sheet.auto_filter.ref = f"A1:{sheet.cell(max(2, sheet.max_row), len(columns)).coordinate}"
    book.save(destination)


def build_week(root: Path, week_start: date) -> tuple[Path, dict]:
    if week_start.weekday() != 0:
        raise ValueError("week_start 必须是周一")
    week_end = week_start + timedelta(days=6)
    label = f"{week_start.isocalendar().year}-W{week_start.isocalendar().week:02d}"
    parent = root / "weekly"
    parent.mkdir(parents=True, exist_ok=True)
    target = parent / f"{label}_{week_start}_{week_end}"
    staged = Path(tempfile.mkdtemp(prefix=".weekly-", dir=parent))
    policies: list[dict] = []
    attachments: list[dict] = []
    statuses: list[dict] = []
    missing_days: list[str] = []
    missing_files: list[str] = []
    seen: set[tuple[str, str, str]] = set()
    try:
        (staged / "附件").mkdir()
        (staged / "附件" / "说明.txt").write_text(
            "按渠道和公告名称存放上周已下载的附件，详见周汇总表的附件明细。\n", encoding="utf-8"
        )
        for offset in range(7):
            day = week_start + timedelta(days=offset)
            day_root = root / day.isoformat()
            listing = day_root / "国内政策清单.csv"
            index_path = day_root / "索引.json"
            if not listing.is_file() or not index_path.is_file():
                missing_days.append(day.isoformat())
                statuses.append({"归档日期": datetime.combine(day, datetime.min.time()), "渠道": "全部",
                                 "抓取状态": "daily_missing", "错误说明": "每日清单或索引缺失"})
                continue
            index = json.loads(index_path.read_text(encoding="utf-8"))
            if index.get("target_date") != day.isoformat():
                raise ValueError(f"每日索引日期不符: {index_path}")
            indexed = {}
            for source in index.get("sources", []):
                statuses.append({"归档日期": datetime.combine(day, datetime.min.time()),
                                 "渠道": source["source"], "抓取状态": source["status"],
                                 "列表条目数": source.get("listing_count", ""),
                                 "公告数": len(source.get("notices", [])),
                                 "错误说明": "; ".join(source.get("errors", []))})
                for notice in source.get("notices", []):
                    indexed[(source["source"], notice["url"])] = notice
            with listing.open(encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                required = {"渠道", "发布日期", "公告名称", "原文链接", "正文状态", "公告文件夹"}
                if not required.issubset(reader.fieldnames or []):
                    raise ValueError(f"每日清单缺少必要列: {listing}")
                daily_keys = set()
                for row in reader:
                    if row["发布日期"] != day.isoformat():
                        raise ValueError(f"每日清单日期不符: {listing}")
                    daily_keys.add((row["渠道"], row["原文链接"]))
                    key = (row["发布日期"], row["渠道"], row["原文链接"])
                    if key in seen:
                        continue
                    seen.add(key)
                    notice = indexed.get((row["渠道"], row["原文链接"]))
                    if notice is None:
                        raise ValueError(f"每日清单与索引不一致: {key}")
                    folder = Path(row["公告文件夹"])
                    if folder.is_absolute() or ".." in folder.parts or not folder.parts or folder.parts[0] != "政策正文":
                        raise ValueError(f"非法公告目录: {folder}")
                    weekly_folder = Path("政策正文") / day.isoformat() / Path(*folder.parts[1:])
                    weekly_attachment_folder = Path("附件") / day.isoformat() / Path(*folder.parts[1:])
                    text_relative = folder / "正文.txt"
                    if (day_root / text_relative).exists():
                        original = checked_file(day_root, str(text_relative), "政策正文")
                        copied = staged / weekly_folder / "正文.txt"
                        copied.parent.mkdir(parents=True, exist_ok=True)
                        shutil.copy2(original, copied)
                    elif row["正文状态"] == "full_text":
                        missing_files.append(f"{day}/{text_relative}")
                    files = []
                    weekly_notice = dict(notice)
                    weekly_notice["folder"] = str(weekly_folder)
                    weekly_notice["attachment_folder"] = str(weekly_attachment_folder)
                    weekly_notice["attachments"] = [dict(item) for item in notice.get("attachments", [])]
                    for item in weekly_notice["attachments"]:
                        relative = item.get("file", "") if item.get("status") == "downloaded" else ""
                        state = item.get("status", "skipped")
                        if state == "downloaded" and not relative:
                            missing_files.append(f"{day}/{row['公告名称']}: 附件路径缺失")
                            state = "file_missing"
                            item["status"] = state
                            item["reason"] = "每日索引未记录附件路径"
                        if relative:
                            try:
                                original = checked_file(day_root, relative, "附件")
                                source_relative = Path(relative)
                                weekly_relative = Path("附件") / day.isoformat() / Path(*source_relative.parts[1:])
                                copied = staged / weekly_relative
                                copied.parent.mkdir(parents=True, exist_ok=True)
                                if copied.exists() and copied.read_bytes() != original.read_bytes():
                                    raise ValueError(f"同名附件内容不一致: {weekly_relative}")
                                if not copied.exists():
                                    shutil.copy2(original, copied)
                                relative = str(weekly_relative)
                                item["file"] = relative
                                files.append(relative)
                            except FileNotFoundError:
                                missing_files.append(f"{day}/{item['file']}")
                                state = "file_missing"
                                item["status"] = state
                                item["file"] = ""
                                item["reason"] = "每日归档附件文件缺失"
                        attachments.append({"公告日期": datetime.combine(day, datetime.min.time()),
                                            "渠道": row["渠道"], "公告名称": row["公告名称"],
                                            "附件名称": Path(relative).name if relative else item.get("label", ""),
                                            "周归档路径": relative if state == "downloaded" else "",
                                            "原始地址": item.get("url", ""),
                                            "文件大小（字节）": item.get("bytes", "") if state == "downloaded" else "",
                                            "下载状态": state, "说明": item.get("reason", "")})
                    metadata = staged / weekly_folder / "公告.json"
                    metadata.parent.mkdir(parents=True, exist_ok=True)
                    metadata.write_text(json.dumps(weekly_notice, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                    policies.append({"公告日期": datetime.combine(day, datetime.min.time()),
                                     "渠道": row["渠道"], "公告名称": row["公告名称"],
                                     "原文链接": row["原文链接"], "正文状态": row["正文状态"],
                                     "正文文件": str(weekly_folder / "正文.txt") if (staged / weekly_folder / "正文.txt").is_file() else "",
                                     "附件数量": len(files), "附件文件夹": str(weekly_attachment_folder),
                                     "附件下载": "; ".join(files),
                                     **{key: row.get(key, "") for key in POLICY_CLASSIFICATION_COLUMNS}})
                if daily_keys != set(indexed):
                    raise ValueError(f"每日清单与索引条目不一致: {listing}")
        save_workbook(staged / "国内政策周汇总.xlsx", policies, attachments, statuses)
        summary = {"week_start": week_start.isoformat(), "week_end": week_end.isoformat(),
                   "generated_at": datetime.now(TIMEZONE).isoformat(), "policy_count": len(policies),
                   "attachment_count": sum(1 for a in attachments if a["下载状态"] == "downloaded"),
                   "missing_days": missing_days, "missing_files": missing_files,
                   "source_failures": [{"date": s["归档日期"].date().isoformat(), "source": s["渠道"],
                                        "status": s["抓取状态"], "error": s["错误说明"]}
                                       for s in statuses if s["抓取状态"] != "ok"]}
        (staged / "汇总索引.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if target.exists():
            backup = parent / (target.name + ".backup")
            if backup.exists():
                raise FileExistsError(f"待处理的旧备份仍存在: {backup}")
            target.rename(backup)
            try:
                staged.rename(target)
            except Exception:
                backup.rename(target)
                raise
            shutil.rmtree(backup)
        else:
            staged.rename(target)
        return target, summary
    finally:
        if staged.exists():
            shutil.rmtree(staged)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="汇总上周国内政策公告、正文和附件")
    parser.add_argument("--output-dir", type=Path, default=Path("data/government_announcements"))
    parser.add_argument("--week-start", type=date.fromisoformat,
                        default=previous_week(datetime.now(TIMEZONE).date()))
    args = parser.parse_args(argv)
    target, summary = build_week(args.output_dir, args.week_start)
    print(target / "国内政策周汇总.xlsx")
    return 1 if summary["missing_days"] or summary["missing_files"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
